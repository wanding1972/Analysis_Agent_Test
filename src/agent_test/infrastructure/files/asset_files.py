"""File adapters for Assets; no Canonical conversion or business inference."""

import json
import logging
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import yaml
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from agent_test.domains.assets.api import AssetReadError, GoldenQuestion


logger = logging.getLogger(__name__)


def _unique_pairs(pairs):
    result = {}
    for key, value in pairs:
        if not isinstance(key, str) or key in result:
            raise ValueError("document keys must be unique strings")
        result[key] = value
    return result


class _AssetYamlLoader(yaml.SafeLoader):
    pass


def _yaml_mapping(loader, node):
    return _unique_pairs(loader.construct_pairs(node, deep=True))


_AssetYamlLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, _yaml_mapping)


def _invalid_constant(value):
    raise ValueError("nonfinite JSON constants are not permitted")


def _object(payload):
    if not isinstance(payload, dict):
        raise ValueError("asset document must be an object")
    return payload


class JsonAssetSetSource:
    def __init__(self, path: Path):
        self.path = path

    def read(self) -> dict[str, object]:
        try:
            with self.path.open(encoding="utf-8") as stream:
                return _object(json.load(stream, object_pairs_hook=_unique_pairs,
                                         parse_constant=_invalid_constant))
        except (OSError, UnicodeError, ValueError, RecursionError) as error:
            logger.error("Unable to read JSON asset document")
            raise AssetReadError("Unable to read JSON asset document") from error


class YamlAssetSetSource:
    def __init__(self, path: Path):
        self.path = path

    def read(self) -> dict[str, object]:
        try:
            with self.path.open(encoding="utf-8") as stream:
                return _object(yaml.load(stream, Loader=_AssetYamlLoader))
        except (OSError, UnicodeError, ValueError, yaml.YAMLError, RecursionError) as error:
            logger.error("Unable to read YAML asset document")
            raise AssetReadError("Unable to read YAML asset document") from error


class ExcelQuestionSource:
    """Strict PRD headers; no formula evaluation or guessed sheet selection."""

    def __init__(self, path: Path, sheet_name: str | None = None):
        self.path = path
        self.sheet_name = sheet_name

    def read(self) -> list[dict[str, object]]:
        try:
            workbook = load_workbook(self.path, read_only=True, data_only=False, keep_links=False)
            try:
                if self.sheet_name is None and len(workbook.sheetnames) != 1:
                    raise ValueError("multi-sheet workbooks require an explicit sheet name")
                name = self.sheet_name if self.sheet_name is not None else workbook.sheetnames[0]
                sheet = workbook[name]
                rows = sheet.iter_rows()
                header_cells = next(rows, ())
                headers = [cell.value for cell in header_cells]
                if (any(not isinstance(value, str) for value in headers)
                        or len(headers) != len(set(headers))
                        or set(headers) != set(GoldenQuestion.model_fields)):
                    raise ValueError("question headers must match the question schema")
                result = []
                for cells in rows:
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValueError("formula and error cells are not accepted")
                    values = [cell.value for cell in cells]
                    if all(value is None for value in values):
                        continue
                    result.append(dict(zip(headers, values, strict=True)))
                return result
            finally:
                workbook.close()
        except (OSError, ValueError, KeyError, UnicodeError, BadZipFile,
                InvalidFileException, ParseError) as error:
            logger.error("Unable to read Excel question document")
            raise AssetReadError("Unable to read Excel question document") from error
